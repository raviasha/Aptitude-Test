from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.55)
sec.left_margin = sec.right_margin = Inches(.65)
sec.footer_distance = Inches(.25)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.03
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)
doc.core_properties.title = 'KSAT Quick Installation Guide'
doc.core_properties.author = 'KSAT'
doc.core_properties.subject = 'Essential setup for departmental faculty servers and student clients'
footer = sec.footer.paragraphs[0]
footer.text = 'KSAT 2.0.0  |  10 September 2026                                            '
footer.runs[0].font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text='', bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def h(text):
    doc.add_heading(text, 1)

def step(number, text):
    para = p()
    para.paragraph_format.left_indent = Inches(.18)
    para.paragraph_format.first_line_indent = Inches(-.18)
    para.add_run(f'{number}. ').bold = True
    para.add_run(text)

def code(text):
    para = p()
    para.paragraph_format.left_indent = Inches(.18)
    para.paragraph_format.space_after = Pt(5)
    run = para.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.2)

doc.add_paragraph('KSAT Quick Installation Guide', 'Title')
p('Faculty server and student PCs  |  TEST BUILD - NOT FOR PRODUCTION', True)
p('Use the installers and question banks supplied in the Google Drive folder. Set up one server and two student PCs first; install the rest after the pilot works.')

h('1 Prepare the lab network')
step(1, 'Choose a reliable Windows 64-bit PC as the faculty server. Connect it and the student PCs to the lab network, preferably by Ethernet. Keep the server powered on and prevent sleep during exams. You need administrator access for installation.')
step(2, 'On the server, open Command Prompt and run hostname and ipconfig. Note its name and IPv4 address. Ask IT to keep the address stable and provide a server name that works from every student PC. Check that all PCs have the correct date and time.')
step(3, 'Before installing, open Command Prompt on TWO student PCs and ping the server by address and by name. Replace these examples with your server details:')
code('ping -4 192.168.10.10\nping -4 ksat-server.example.edu')
p('Both PCs should receive replies from the correct server. If either test fails, ask IT to fix the connection or name. If the network blocks ping, ask IT to confirm connectivity before continuing.')

h('2 Download the installation files')
p('Download KSATCoordinatorSetup-2.0.0.exe, KSATClientSetup-2.0.0.exe, and the supplied question banks to a local folder. Extract the outer download ZIP if there is one. Keep each individual question-bank ZIP intact. Install the Coordinator on the server and the Client on student PCs.')

h('3 Install and start the faculty server')
step(1, 'Right-click KSATCoordinatorSetup-2.0.0.exe and choose Run as administrator. Keep the default installation folder. For HTTPS hostname, enter the server name you tested above, without https:// or a port. Keep HTTPS port 8443. Finish Setup and launch KSAT Faculty Coordinator.')
step(2, 'On the server, open PowerShell as administrator and run the following once to let its browser trust the server certificate:')
code('certutil -addstore Root `\n  "C:\\ProgramData\\KSAT Coordinator\\public\\coordinator-ca.pem"')
step(3, 'Reopen the browser at https://127.0.0.1:8443. Select Faculty and your Department, then sign in with username faculty and password faculty123 for a new installation. Keep the Coordinator running.')
step(4, 'On both pilot student PCs, open PowerShell and run the command below using your server name. Check for TcpTestSucceeded : True. If it is False, ask IT to allow the coordinator through the lab firewall on TCP 8443.')
code('Test-NetConnection ksat-server.example.edu -Port 8443')

doc.add_page_break()
h('4 Install two student PCs first')
step(1, 'On the server, open the folder below in File Explorer. Copy its two public files, coordinator-ca.pem and coordinator-public.json, to each pilot PC. Use the files from this department’s server; keep them unchanged.')
code('C:\\ProgramData\\KSAT Coordinator\\public')
step(2, 'On each pilot PC, right-click KSATClientSetup-2.0.0.exe and choose Run as administrator. Keep the default folder. Enter the full Coordinator address using the same tested name and port:')
code('https://ksat-server.example.edu:8443')
step(3, 'At Coordinator public trust, select coordinator-ca.pem as the CA file and coordinator-public.json as the metadata file. Finish Setup and open the KSAT Lab Client shortcut. Its local page is http://127.0.0.1:8010. Device registration is automatic; no enrollment code is needed.')
step(4, 'On Faculty → Tests, confirm that both PCs appear under Managed lab computers. On each client, use New student? Create account, enter the student details and a password of at least six characters, then sign in. Use different student accounts for the two pilot PCs.')

h('5 Import the question banks on the server')
p('Sign in as Faculty and open Question banks. Under Upload a package, select one original question-bank ZIP and choose Upload v2 package. Wait for success and confirm that the bank and its questions appear. Repeat for each bank. If supplied an HTML/JSON pair instead, use Upload a pair and select both matching files.')
p('Keep the Google Drive folder and question-bank source files with faculty. Students receive assessment content through the application.')

h('6 Run a short pilot and finish the lab setup')
step(1, 'In Faculty → Tests, create a short test from an imported bank and select Launch. Students have ten minutes to start. On both pilot PCs, start the assessment and enter full screen when prompted. Confirm that questions, images, answers, navigation, and the timer work.')
step(2, 'Use Extend on Faculty once and check that both student timers and Faculty’s Exam time left increase. Submit both attempts and confirm their results on Faculty. Select Close when finished, then check that students can review answers.')
step(3, 'After the pilot succeeds, repeat the client installation on the remaining PCs using the same server address and two public files. Give each student a separate account. Keep the same server for the whole department.')

h('For every assessment')
p('Start KSAT Faculty Coordinator after a server restart and keep the server awake. Students open KSAT Lab Client and stay on the same PC until submission is confirmed. Close a finished test only when you are ready to release its answers.')
p('If a setup step fails, note the message and contact your lab IT person or the sender. Include the server name and the PC where it failed.')

path = out / 'KSAT_Quick_Installation_Guide.docx'
for root in (doc.styles.element, doc.element):
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(path)
print(path)
