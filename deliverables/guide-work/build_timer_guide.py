from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parents[1]
NAME = 'KSAT_Department_Installation_Guide_2026-09-10_Timer_Fix'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].paragraph_format.space_after = Pt(14)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.core_properties.title = 'KSAT Department Installation Guide'
doc.core_properties.subject = 'KSAT 2.0 client and server test distribution dated 10 September 2026'
doc.core_properties.author = 'KSAT'
doc.core_properties.keywords = 'KSAT, installation, coordinator, client, test build, departments'
head = sec.header.paragraphs[0]
head.text = 'KSAT 2.0  |  Department installation guide'
head.runs[0].font.size = Pt(9)
foot = sec.footer.paragraphs[0]
foot.text = 'TEST BUILD  •  NOT FOR PRODUCTION                                      '
foot.runs[0].font.size = Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)

def p(text='', bold=False):
    x=doc.add_paragraph(); x.add_run(text).bold=bold; return x
def h(text): doc.add_heading(text,2)
def step(n,text):
    x=p(); x.paragraph_format.left_indent=Inches(.22); x.paragraph_format.first_line_indent=Inches(-.22)
    x.add_run(f'{n}. ').bold=True; x.add_run(text); return x
def bullet(text):
    x=p(); x.paragraph_format.left_indent=Inches(.15); x.paragraph_format.first_line_indent=Inches(-.15)
    x.add_run('• '); x.add_run(text)
def note(label,text):
    x=p(); x.add_run(label+' ').bold=True; x.add_run(text)
def code(text):
    for line in text.splitlines():
        x=p(); x.paragraph_format.space_after=Pt(2); x.paragraph_format.line_spacing=1
        r=x.add_run(line); r.font.name='Consolas'; r.font.size=Pt(8.6)
def page(title):
    x=doc.add_heading(title,1); x.paragraph_format.page_break_before=True
def link(label,url):
    x=p(); a=OxmlElement('w:hyperlink'); a.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'000000'); pr.append(c)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr)
    t=OxmlElement('w:t'); t.text=label; r.append(t);a.append(r);x._p.append(a)
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for vals in rows:
        cells=t.add_row().cells
        for i,v in enumerate(vals):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); avoid=OxmlElement('w:cantSplit');trpr.append(avoid)
        if ri==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9E9E9');cell._tc.get_or_add_tcPr().append(sh)
            for pa in cell.paragraphs:
                pa.paragraph_format.space_after=Pt(2);pa.paragraph_format.line_spacing=1.04
                for r in pa.runs:r.font.size=Pt(9.4);r.bold=(ri==0)
    p().paragraph_format.space_after=Pt(0)

doc.add_paragraph('KSAT Department Installation Guide', 'Title')
doc.add_paragraph('Windows server and student client setup', 'Subtitle')
p('KSAT 2.0.0  |  Integrity and faculty timer test build  |  10 September 2026',True)
p('Use this guide when a department receives KSAT installers and question banks through a Google Drive shared link. It takes the lab administrator from selecting a server and testing the network to installing clients, importing banks, running a pilot, and handing over the lab.')
note('Release status','These binaries are test signed and marked NOT FOR PRODUCTION. Use an institution-approved test lab and disposable assessments. Completing this guide does not convert this build into a production release.')
h('Installation order')
p('Choose one server → verify connectivity from two student PCs → download and verify files → install and initialize the coordinator → distribute its public trust files → install two pilot clients → import banks → pass the pilot → install the remaining clients.')
h('How the computers connect')
table(['Computer','Install and open','Connection'],[
('Faculty server','KSAT Faculty Coordinator\nLocal faculty page: https://127.0.0.1:8443','Receives client traffic over private LAN HTTPS, normally TCP 8443.'),
('Every student PC','KSAT Lab Client\nStudent page: http://127.0.0.1:8010','Local Windows service connects to the coordinator. Port 8010 stays local.')],[1.1,2.9,2.87])
p('This is the installed client and server edition. Students open their local client shortcut. Instructions for the old Aptitude Lab browser-only deployment on port 8000 do not apply.')
h('Guide map')
table(['Pages','Task'],[
('2–4','Choose the server and prove network connectivity before installation'),
('5–8','Download files and install the server and pilot clients'),
('9–11','Import question banks and test the complete assessment workflow'),
('12–14','Roll out, operate, back up, and update the lab'),
('15–16','Troubleshoot network, installer, and assessment problems'),
('17–18','Verify file hashes and complete the handover record')],[.8,6.07])

page('1 Choose the server and prepare the lab')
p('The server is an ordinary institution-managed Windows computer that runs KSAT Faculty Coordinator. A dedicated server operating system is not required. Choose one server for this department’s test deployment and give it a named IT owner.')
step(1,'Choose a reliable x64-compatible Windows PC with an operating system supported by your institution, administrator access, a healthy local disk, and a stable wired Ethernet connection. Use an up-to-date Edge or Chrome browser on the server and clients.')
step(2,'Prefer a modern four-core processor, 8 GB RAM, an SSD, and at least 10 GB free for the initial pilot, with extra space for banks, attempts, and backups. These are planning recommendations, not measured minimum requirements or a capacity guarantee. Validate the intended class size on your actual hardware.')
step(3,'Keep the server in staff custody. Avoid a student workstation, an intermittently used laptop, a mobile hotspot, or a guest Wi-Fi network. Use a UPS where available. Connect clients to the same permitted lab LAN or an IT-approved routed network.')
step(4,'Ask IT to reserve the server IPv4 address in DHCP, or assign a conflict-free static address with the correct subnet mask, gateway, and DNS servers. Do not select an unused-looking address yourself. Give every computer a unique Windows computer name.')
step(5,'Agree on a stable DNS hostname for the server. It must resolve to this server from every client. The installer’s suggested name ending in .local is only a default; test it instead of assuming it works. An institutional DNS record is preferable.')
step(6,'On mains power, prevent the server from sleeping, hibernating, or restarting for updates during a session. Configure clients to remain awake during assessments. Set correct date, time, and time zone on all PCs and enable institution-approved time synchronization.')
step(7,'Schedule installation outside exams. Arrange administrator credentials for installation, an ordinary student Windows account for testing, a protected backup destination, and a contact for firewall/DNS changes. Normal student use should not require administrator rights.')
h('Capacity and boundaries')
p('The project runbook targets one coordinator and up to 100 installed clients. Start with two clients and expand only after your lab passes the pilot and an appropriate concurrent-use test. An ordinary user must not be able to edit the client’s protected state or signing identity.')
p('Clients register automatically, without an enrollment code. IT must limit access to the coordinator port to managed lab computers. Do not port-forward it to the Internet or make it reachable from unmanaged student devices. The download link is for software distribution; it is not the exam server.')
note('What is already bundled','Recipients do not need Python, Node.js, a database server, or development tools to install these packages. Google Drive access is needed for downloads; assessment traffic uses the lab network.')

page('2 Record the network before installing')
p('Complete this worksheet with IT. Example values below illustrate commands only. Replace them with your lab’s actual values throughout this guide.')
table(['Setting','Example','Actual value'],[
('Department and room','CSE Lab 1','________________________'),
('Server Windows name','KSAT-SERVER','________________________'),
('Stable DNS hostname','ksat-server.example.edu','________________________'),
('Reserved server IPv4','192.168.10.10','________________________'),
('Subnet and gateway','Provided by IT','________________________'),
('HTTPS TCP port','8443','________________________'),
('Full coordinator URL','https://ksat-server.example.edu:8443','________________________'),
('Network profile','Private or managed Domain','________________________'),
('Pilot client names','LAB-PC01 and LAB-PC02','________________________'),
('IT and faculty contacts','Name and phone','________________________'),
('Protected backup location','Approved offline storage','________________________')],[1.65,2.67,2.55])
h('On the proposed server')
step(1,'Connect the Ethernet cable and confirm the connection works. Open Start, type PowerShell, and open Windows PowerShell. Run these read-only commands:')
code('hostname\nipconfig /all\nGet-NetConnectionProfile\nGet-Date')
step(2,'Identify the active Ethernet adapter. Record its IPv4 address, subnet mask, default gateway, and DNS servers. Ignore disconnected, VPN, virtual-machine, and loopback adapters unless IT explicitly uses one for the lab.')
step(3,'An address beginning 169.254 usually means the PC has not obtained a usable DHCP configuration. Ask IT to fix this before proceeding. A successful Internet connection by itself does not prove the lab clients can reach the server.')
step(4,'Resolve the agreed hostname and check that it maps to the intended server address. Repeat after a reboot or DHCP reservation change.')
code('Resolve-DnsName ksat-server.example.edu')
note('If DNS is unavailable','Have IT provide a centrally managed hosts-file mapping on the server and every client, then verify the same name with ping. The hosts file is C:\Windows\System32\drivers\etc\hosts. Do not rename it to hosts.txt. A hosts mapping does not change the TLS hostname; keep using the agreed hostname in the installer.')

page('3 Ping from two clients before installation')
p('Do this before running either installer. Choose two student PCs, preferably on different lab switch branches or rows. If multiple rooms or network segments will be used, include a representative PC from each segment as well.')
step(1,'On pilot client 1, open PowerShell or Command Prompt. Record its computer name and address, then ping the server by IPv4 and by the exact hostname you will enter in the coordinator installer.')
code('hostname\nipconfig\nping -4 -n 4 192.168.10.10\nping -4 -n 4 ksat-server.example.edu')
step(2,'Check that the hostname resolves to the recorded server IPv4 address. A passing test shows replies from that address and 0% loss. Save the output or record the results below. Repeat the same commands on pilot client 2.')
step(3,'If either check fails, fix it before installation. If institutional policy blocks ICMP ping, IT must confirm and document the block, verify addressing/name resolution/routing, and arrange the post-install TCP test on page 7. Do not treat a timeout as proof the server is offline or silently mark it as a pass.')
table(['Result','What to check next'],[
('IPv4 works but hostname fails','DNS or hosts mapping. Correct the name before generating coordinator certificates.'),
('Both time out','Power, cable, active adapter, subnet/gateway, VLAN routing, Wi-Fi client isolation, VPN route, and server ICMP policy.'),
('Destination host unreachable','Often a routing or local link problem. Inspect the address that reports the error; it is not a successful server reply.'),
('Hostname resolves to another PC','Fix stale DNS/hosts entries or an address conflict. Do not install against the wrong machine.'),
('Intermittent replies or packet loss','Resolve cable, switch, wireless, or addressing problems and repeat the test.')],[2,4.87])
p('IT may permit ICMPv4 echo requests from the approved lab subnet using Windows Defender Firewall with Advanced Security. Keep the firewall enabled. Do not enable broad file sharing or disable security software just to make ping pass.')
table(['Client and time','Ping IPv4','Ping exact hostname','Verified by'],[
('__________________','Pass / Fail / IT exception','Pass / Fail / IT exception','____________'),
('__________________','Pass / Fail / IT exception','Pass / Fail / IT exception','____________')],[1.55,1.8,2.1,1.42])
note('Before continuing','Both clients must have verified network reachability and correct hostname resolution, or a documented IT-confirmed ICMP restriction with routing verified. TCP 8443 is normally closed before the coordinator is installed and running; test that port after installation.')
link('Windows command reference for ping','https://learn.microsoft.com/en-us/windows-server/administration/windows-commands/ping')

page('4 Download and verify the distribution')
h('For the person sharing the Google Drive folder')
step(1,'Use a dated folder such as KSAT 2.0 Test Build 2026-09-10. Include this guide, release notes marked NOT FOR PRODUCTION, both Setup executables, SHA256SUMS.txt, and the approved question banks. Keep earlier builds in clearly separate folders.')
step(2,'Share with named departmental faculty and IT recipients, using Viewer access with downloads allowed under institutional policy. Question-bank source files contain correct answers and explanations; do not give students the master distribution link.')
step(3,'Include a bank inventory giving each filename, bank name, revision, expected question count, and any known limitations. Provide the installer checksum record through a trusted channel. Each department generates its own server trust files after installation; do not place another department’s trust bundle in the generic install package.')
h('For the receiving department')
step(4,'Open the supplied link using the authorized Google account. Download both installers, the checksum file, release notes, this guide, and the chosen banks to a local folder such as C:\KSAT-Install. If access or download is blocked, ask the sender or IT to correct access; do not substitute files from another source.')
step(5,'Wait for downloads to finish. If Drive provides an outer ZIP for the download, right-click it and choose Extract All. Work from the extracted local folder. Preserve individual v2 question-bank ZIPs inside it for import; do not extract or repackage those ZIPs.')
step(6,'In File Explorer enable file-name extensions. Confirm that the server installer is KSATCoordinatorSetup-2.0.0.exe and the student installer is KSATClientSetup-2.0.0.exe. Use these Setup files for first installation, not the standalone KSATCoordinator-2.0.0.exe or KSATClient-2.0.0.exe.')
step(7,'Open PowerShell and calculate the two installer hashes. Compare all 64 characters with page 17 and the sender’s trusted checksum record. Uppercase and lowercase hexadecimal letters are equivalent.')
code("Get-FileHash 'C:\KSAT-Install\KSATCoordinatorSetup-2.0.0.exe' -Algorithm SHA256\nGet-FileHash 'C:\KSAT-Install\KSATClientSetup-2.0.0.exe' -Algorithm SHA256")
step(8,'If a hash differs, stop and download the correct dated build again. The displayed version remains 2.0.0 across several test updates, so the version number alone is insufficient. Keep the verified installers for rollback and reproducibility.')
note('Windows protection prompts','This release uses a test signing identity and may show an unrecognized-publisher or SmartScreen warning. After verifying provenance and hashes, proceed only if institutional test policy permits it. If blocked by policy, contact IT. Do not disable Defender, SmartScreen, or certificate checks globally.')
link('Google Drive download instructions','https://support.google.com/drive/answer/2423534?hl=en')

page('5 Install and start the coordinator')
p('Perform these steps on the chosen server only, using the recorded hostname and port. For an existing installation, read the upgrade procedure on page 14 first.')
step(1,'Sign in with an authorized administrator account. Close any existing KSAT coordinator after confirming there are no active exams or pending uploads. Right-click KSATCoordinatorSetup-2.0.0.exe and select Run as administrator. Accept the UAC prompt for the verified file.')
step(2,'Proceed through Setup. Keep the default program folder unless IT has a reason to change it: C:\Program Files\KSAT Coordinator. The application data is stored separately in C:\ProgramData\KSAT Coordinator.')
step(3,'On Coordinator network name, enter the pretested stable DNS name in HTTPS hostname. Enter the hostname only, such as ksat-server.example.edu; do not include https://, a port, a path, or a trailing slash.')
step(4,'On Coordinator HTTPS port, enter 8443 unless IT has approved another port. Use the same port in every later command and client URL. Review the choices and complete installation.')
step(5,'Leave Launch KSAT Faculty Coordinator selected, or open the KSAT Faculty Coordinator desktop or Start menu shortcut afterwards. Approve its administrator prompt. Allow startup to finish. It normally opens the faculty browser at https://127.0.0.1:8443 on the server.')
step(6,'Keep the coordinator process running. This packaged application runs without a console window, and Setup does not create an automatic coordinator Windows service. After each server reboot, an authorized operator must start the shortcut again. Opening only an old browser tab does not start the server.')
step(7,'On first successful startup, confirm that the public folder contains both files below. ProgramData is hidden by default; paste the full folder path in File Explorer’s address bar. If the files are missing, resolve the startup failure before installing clients.')
code('C:\ProgramData\KSAT Coordinator\public\coordinator-ca.pem\nC:\ProgramData\KSAT Coordinator\public\coordinator-public.json')
h('Trust the server certificate on the faculty machine')
p('The coordinator creates its own local certificate authority. The coordinator installer does not import that CA into the server’s Windows trust store. A browser certificate warning can therefore occur on first use.')
step(8,'IT should verify that coordinator-ca.pem was generated by this server, then open PowerShell as administrator on the server and import this exact public CA into the machine Trusted Root store:')
code('certutil -addstore Root "C:\ProgramData\KSAT Coordinator\public\coordinator-ca.pem"')
step(9,'Close and reopen the faculty browser and load https://127.0.0.1:8443 again. The page should load without a certificate warning. If it does not, check time, trust policy, and the certificate instead of clicking through the warning.')
note('Two different certificates','The server’s coordinator CA secures lab HTTPS. It is not the test code-signing certificate used on the downloaded executables. Import only the server CA described above for this step.')

page('6 Verify the server and export public trust')
step(1,'On the faculty sign-in page select Faculty. Select the appropriate Department, enter username faculty and password faculty123 for a new empty installation, and select Sign in. An existing installation retains its accounts; use its existing credentials.')
note('Test account limitation','The current build seeds this default faculty account and has no faculty password-change screen. Restrict this build to the approved test lab. Arrange supported account provisioning with the software maintainer before any production use. Do not edit password hashes manually or assume a department selection creates a separate security boundary.')
step(2,'Confirm the faculty dashboard opens. Initial demo students, a starter bank, and a sample test may be present. They are test content, not departmental records. Use clearly labelled pilot accounts and assessments.')
step(3,'On BOTH pilot clients, open Windows PowerShell and test the selected server port while the coordinator is running:')
code('Test-NetConnection ksat-server.example.edu -Port 8443')
p('Require TcpTestSucceeded : True and the intended server address. This checks a TCP connection, not full TLS or client registration. The client installation and sign-in checks on page 8 complete that validation.')
step(4,'On the server, have IT inspect Windows Defender Firewall with Advanced Security → Inbound Rules. Setup creates KSAT Faculty Coordinator 2.0.0 for the installed executable and chosen TCP port on the Private profile. The active lab network must be compatible with that rule.')
p('If the institutional connection is Domain rather than Private, ask IT to create an approved, appropriately scoped supplementary rule for that profile and executable/port. Do not change a managed network to Private just to bypass policy, and do not modify or rename the installer-owned rule: updates validate its exact ownership. IT should limit reachability to managed lab hosts or VLANs.')
step(5,'Check that the server URL in coordinator-public.json matches the approved hostname and port. Copy ONLY coordinator-ca.pem and coordinator-public.json to the department’s protected installation folder or trusted transfer medium. Do not edit either file.')
step(6,'Compute SHA256 hashes of the two public files on the server. Record them in the department handover record, then compare hashes of the copied files before client installation. Each department must use the pair from its own selected server.')
code('Get-FileHash "C:\ProgramData\KSAT Coordinator\public\*" -Algorithm SHA256')
step(7,'Before enrolling any client, use Stop server and back up the complete C:\ProgramData\KSAT Coordinator folder using page 13. Restart the coordinator and repeat the port check. This preserves the keys to which clients will be bound.')
note('Never distribute','Do not place secrets, the database, client identities, session tokens, or assessment packs in the Google Drive distribution folder. The two public files are shareable, but their authenticity is essential.')
link('Windows PowerShell TCP connection test reference','https://learn.microsoft.com/en-us/powershell/module/nettcpip/test-netconnection')

page('7 Install two pilot clients')
p('Install on the two PCs already used for network checks. Keep the server running. Use the same verified client Setup file and this department’s two public trust files on every PC.')
step(1,'Copy the files to a local folder, for example C:\KSAT-Install. Recheck their hashes after transfer. Sign in as an authorized administrator, right-click KSATClientSetup-2.0.0.exe, choose Run as administrator, and accept the verified UAC prompt.')
step(2,'Keep the default program location C:\Program Files\KSAT Client unless IT requires another path. On Coordinator address enter the full HTTPS URL, for example https://ksat-server.example.edu:8443. It must match the department’s recorded server endpoint.')
step(3,'On Coordinator public trust, browse to coordinator-ca.pem for Coordinator CA file and coordinator-public.json for Coordinator metadata file. Both must be the unmodified matching pair generated by this server. Continue and complete installation.')
step(4,'Setup validates configuration, imports the coordinator CA into the machine trust store, creates the automatic LocalSystem service KSATLabClientAuthority, protects its data folders, and starts the service. It does not open a client inbound firewall port.')
step(5,'Select Open KSAT Lab Client, or use the KSAT Lab Client desktop or Start menu shortcut. Confirm the browser address is http://127.0.0.1:8010. This address refers to the current student PC, not the faculty server.')
step(6,'Open services.msc and find KSAT Lab Client Authority. Verify Status is Running and Startup Type is Automatic. A read-only PowerShell check is also available:')
code('Get-Service KSATLabClientAuthority')
step(7,'Allow automatic device registration to finish. No enrollment code or machine label needs to be typed. On Faculty → Tests, check the Managed lab computers list and verify the Windows computer name and device fingerprint. An active device offers Revoke; a revoked device offers Reactivate.')
step(8,'Repeat on pilot client 2. Return both machines to ordinary student Windows accounts. Verify the shortcut opens successfully without elevation and the two machines appear as distinct devices. Never clone an already initialized client identity/state folder onto another PC.')
h('Create pilot student accounts')
step(9,'On each client select New student? Create account. Enter Student name, official test Student ID / USN, Class, Section, Password, and Confirm password. Replace the defaults AIML and A with the correct values. Passwords must contain at least six characters; use unique test credentials.')
step(10,'Return to sign-in, choose Department, enter Student ID / USN and password, and select Sign in. Use different pilot student accounts on the two PCs. Confirm the assessment list loads.')
note('Keep the shortcut and browser profile','Use one normal browser profile and one client tab. Avoid private/incognito mode or policies that clear site storage during an attempt. Do not clear browser storage, reinstall, or delete ProgramData to resolve an active assessment problem.')

page('8 Import and verify the question banks')
p('Import banks into the coordinator only. Students do not copy question-bank source files onto their PCs. The application distributes the required assessment content through the client workflow.')
h('Recommended v2 ZIP package')
step(1,'On the server sign in as Faculty and open Question banks. Under Upload a package, select the downloaded bank ZIP in Question-bank package, then select Upload v2 package.')
step(2,'Upload the original individual bank ZIP intact. A v2 bank contains manifest.json, question data, and any required assets. Do not upload the outer Google Drive download ZIP, a ZIP of several banks, or a package nested inside an extra folder.')
step(3,'Wait for the success message. Check Available question banks for the expected bank name, question count, format v2, and stimulus count. Compare with the sender’s bank inventory. Repeat for each approved bank and record the counts.')
h('Legacy HTML and JSON pair')
step(4,'If the supplied bank is a matching HTML and JSON pair, use Upload a pair. Select the HTML in Questions and visuals and its corresponding JSON in Choices and answer key. Select Upload & import. Keep the pair from the same revision.')
step(5,'Alternatively, copy both matching-base-name files into the server folder below, then select Refresh folder in Question banks. When the pair is marked Ready, select Import. Copying files by itself does not import a bank.')
code('C:\ProgramData\KSAT Coordinator\Question Banks')
note('Correct data folder','Older bank README files may mention C:\ProgramData\Aptitude Lab. That is the legacy location. For this client/server installation use the KSAT Coordinator folder displayed by the running application.')
h('Validate content before an assessment')
step(6,'Open Faculty → Tests, select the imported bank and desired difficulty, and confirm its categories, chapters, and available quantities. Do not select more questions than are available at that difficulty. The UI supports a total of up to 500 questions.')
step(7,'Run a disposable pilot using representative chapters. Check readable equations, fractions, images, tables, graphs, all answer options, and navigation on the actual client screen. After Faculty closes the pilot, compare sample answers and solution steps with the source bank.')
step(8,'If counts are unexpected, formulas are damaged, or images are missing, retain the original package and exact error. Ask the bank maintainer for a corrected version. Do not relabel an arbitrary ZIP as a v2 bank or repeatedly import the same file.')
p('Prefer the supplied textbook-verified chapter ZIPs when available. The repository also contains older bulk-extracted banks with known mathematical-layout limitations. Format validation does not certify the educational correctness of every answer.')
note('Preserve exam history','Do not delete a bank or test as routine cleanup. Deletion can affect related history and is destructive. Keep pilot records clearly labelled and remove them only through an agreed retention process.')

page('9 Run the first complete assessment')
p('Use disposable accounts and a clearly named assessment such as DEPARTMENT PILOT 2026-09-10. This pilot must pass on both clients before the rest of the lab is installed.')
step(1,'On Faculty → Tests, enter a Test name, select the imported Question bank and Difficulty, enter chapter quantities, review the total, and select Create test. A five-question assessment allows a short functional check; use a separate ten-question assessment for the integrity exercises on page 11.')
step(2,'Verify the test’s content/release status is ready on Faculty. Sign both clients in and wait for the invigilator. A test may not appear in the student list until it is launched. Available assessments refresh about every five seconds and have a Refresh button.')
step(3,'Select Launch on the faculty machine once. This opens a ten-minute window in which students can start. This window is distinct from each student’s exam duration. The assessment duration is one minute per question, starting when that student’s ticket is issued.')
step(4,'On each pilot client choose Start, or Download if content is not cached. In this build, Download also starts the attempt after downloading; it is not a separate preparation-only action. Use the application’s full-screen entry button if prompted. F11 alone is not a substitute. Confirm questions, choices, navigation, and the timer.')
step(5,'Answer questions and verify saved selections. On Faculty select Extend, enter minutes and a reason, and confirm both student timers and Faculty’s Exam time left increase. Reload Faculty and check min added remains. Different student deadlines appear as a range; an individual extension changes only that student’s deadline. Start window remains separate. Leave one question unanswered and use the same PC throughout.')
step(6,'Submit one pilot attempt manually and allow the other to expire in a separate run if necessary. Confirm answers become locked and the upload reaches an accepted, server-confirmed result. Record the attempt and assessment identifiers where shown.')
step(7,'On the faculty dashboard check the submitted results and score. Use Download results CSV and confirm both pilot students, assessment names, scores, and any recorded violations are present. A pending local submission is not yet an accepted faculty result.')
step(8,'Before closing the test, confirm students cannot reveal correct answers or solutions. After all intended attempts have finished, select Close on the faculty test. Check Review answers on each client and verify correct choices, the student’s selections, and explanations.')
step(9,'Use Back to assessments to return without signing out, then test Sign out after the attempt is acknowledged. Create or duplicate a separate pilot assessment to confirm that the next test becomes available and a new account can sign in.')
h('What to record')
table(['Check','Pilot client 1','Pilot client 2'],[
('Unique device and student account','Pass / Fail','Pass / Fail'),
('Download and start during launch window','Pass / Fail','Pass / Fail'),
('Full screen, saved answers and timer extension','Pass / Fail','Pass / Fail'),
('Confirmed result and faculty CSV','Pass / Fail','Pass / Fail'),
('Review blocked before Close and available after','Pass / Fail','Pass / Fail'),
('Sign out and next assessment','Pass / Fail','Pass / Fail')],[4.07,1.4,1.4])
note('Close releases answers','Closing an assessment releases review information and cannot retract answers already seen. Expiry of the start window alone does not release answers. Use a new or duplicated test for the next separate assessment.')

page('10 Test integrity and interrupted connections')
p('Run these exercises on the actual Windows and browser versions used by students, using a disposable timed exam. Record the action and time, return to the exam between actions, and check the final faculty result and CSV after upload. The timer must keep running while the student is away.')
table(['Exercise','Acceptance observation'],[
('Escape from application full screen','Questions are gated; returning requires the full-screen/resume action; fullscreen exit is recorded.'),
('Ctrl+Tab and Ctrl+Shift+Tab','Switch to another browser tab if the browser permits it. A hidden/focus/fullscreen signal is recorded when exposed.'),
('Alt+Tab and minimize then restore','Questions are gated on observable loss and the event reaches the submitted record.'),
('Win+D, Win+Tab, Win+Ctrl+Left or Right','Test desktop/task switching. Record actual detection on this lab configuration, including any missed action.'),
('Win+L then unlock','Test lock/unlock, resume behavior, continuing timer, and recorded signals.'),
('Two rapid departures','Both observed transitions survive; retrying an existing event does not multiply that same event.'),
('Reload and close then reopen','Resume on the same PC with saved answers and the original deadline. Leave the browser closed for at least 16 seconds to test a monitoring gap.'),
('Ordinary sign-in and completed review','No new exam violations should be generated outside an active attempt.')],[2.67,4.2])
h('Test a coordinator network interruption')
step(1,'Start a separate disposable exam and save some answers. Disconnect ONLY that pilot client’s lab network briefly, or have IT temporarily interrupt access to the test coordinator. Keep the client PC, browser, and local authority service running.')
step(2,'Confirm answering and navigation remain responsive, selections persist, and the timer continues. Perform one observable focus or fullscreen departure and return. Submit while disconnected and confirm the client shows a saved or pending-upload state.')
step(3,'Restore connectivity to the same coordinator. Wait for upload confirmation. Verify one accepted result, the saved answers, and the integrity records on Faculty and in CSV. Do not shut down, clear data, or move the attempt while upload is pending.')
step(4,'On a separate test run, restart the coordinator between attempts and reboot a pilot client. Verify the server is started manually, the client authority starts automatically, and the saved configuration and device identity still work.')
note('Detection limits','Browser signals cannot guarantee detection of every Windows virtual-desktop transition. A monitoring gap indicates interrupted reporting, not proof of a particular action. Multiple signals may describe one departure. Very short, unreported interruptions near expiry may leave no event. Treat unexplained behavior as a pilot finding; do not claim complete operating-system lockdown.')
note('If every switch must be guaranteed','The present test build needs additional managed kiosk controls or a native interactive-session monitor for that requirement. Do not mark a missed desktop-switch test as passed or use absence of a flag as proof that no departure occurred.')

page('11 Roll out and run a normal session')
h('After the two client pilot passes')
step(1,'Resolve every failed check, then install the remaining managed PCs using the exact verified Setup file, coordinator URL, and public trust pair. Keep the same configuration and browser policy across the lab.')
step(2,'For each PC, record its Windows name, client build checksum, device identity/fingerprint, service status, registration result, and a successful student sign-in. Spot-check network connectivity across all network segments, and confirm every planned seat can connect.')
step(3,'Run a concurrent rehearsal at the planned class size using test accounts. Check downloads, near-simultaneous starts, autosave responsiveness, submissions, faculty queue/results, and recovery from a planned outage. Record results and remaining limitations. Keep the release marked NOT FOR PRODUCTION.')
h('Before each approved test session')
bullet('Start the server early and open the faculty dashboard. Confirm clock, free disk space, stable hostname/IP, power, and network. Check a complete recent backup exists. Avoid Windows updates, restarts, installs, or network changes during the session.')
bullet('On representative clients, verify the authority service is running, the local shortcut works, registration is active, and the server is reachable. Resolve pending uploads from an earlier session before reusing that seat.')
bullet('Verify the intended question bank, chapter quantities, difficulty, and resulting duration. Create a fresh assessment or use Duplicate for a new release; do not change the meaning of an assessment after students have started.')
bullet('Have students sign in using their own accounts and wait for the invigilator before selecting Start or Download. Download also starts an attempt after fetching content. Close unrelated applications and tabs. Explain that leaving exam focus/full screen is recorded and the clock continues.')
h('During and after the session')
step(4,'Launch once when the room is ready. Students must start within the displayed ten-minute start window. Do not confuse that window with the individual assessment clock. Verify starts and submissions on the faculty screen.')
step(5,'If a student needs help, record the PC, student, time, and attempt state. Keep the student on the same machine where possible. Do not copy client databases or ask the student to clear storage. An interrupted client resumes with its original deadline.')
step(6,'Wait for acknowledged submissions and check the faculty queue and results. If an upload is pending, keep the original client and the coordinator powered and connected until confirmation or a documented intervention.')
step(7,'When all intended attempts are finished and faculty is ready to release answers, select Close. Export results CSV, review the integrity signals with their context, sign students out after confirmation, and perform the protected backup procedure.')
note('School policy','An integrity signal is evidence for review; it does not establish motive. OS dialogs, lock events, connectivity, and browser behavior can contribute. Document pilot findings and tell students the institution’s applicable rules before testing.')

page('12 Shut down and back up safely')
h('Stop the coordinator before maintenance')
step(1,'Finish the session, check all submissions are acknowledged or formally recorded for intervention, and confirm no active attempts remain. Export results CSV for departmental records.')
step(2,'In the faculty page header select Stop server. Confirm the message that all connected users will be disconnected. Wait for Server stopped and for the process to exit. Closing the browser or choosing Sign out alone does not stop the server. Avoid Task Manager termination or powering off while it writes data.')
step(3,'If needed, verify the coordinator process has exited in Task Manager or with Get-Process KSATCoordinator -ErrorAction SilentlyContinue. Do not start a second coordinator on the same data folder.')
h('Make a complete backup')
step(4,'With the coordinator stopped, copy the entire C:\ProgramData\KSAT Coordinator folder to a dated, protected backup destination. Use File Explorer as an authorized administrator, show hidden items or paste the path, and confirm the copy completes without skipped files.')
step(5,'Include the database and any accompanying SQLite files, runtime settings, public folder, secrets, Assessment Releases, Question Assets, Question Banks, and all other contents. Keep this as one matching snapshot. Protect the backup as strongly as the server itself.')
step(6,'Record the date, build checksum, hostname, port, backup location, and operator. Retain more than one dated backup, with at least one offline or separately protected copy. Back up before client enrollment, after completed sessions, and before upgrades or certificate work.')
note('The dashboard database download is insufficient','Download database backup provides a database copy. It does not replace the complete stopped-folder backup needed to preserve signing keys, trust, release artifacts, and recovery capability.')
h('Client records that are still pending')
p('The client service runs independently of the browser. Keep the PC running until uploads are confirmed. If IT must preserve a failed client, stop KSATLabClientAuthority only as part of controlled maintenance and copy its complete C:\ProgramData\KSAT Client folder, including identity, state, packs, configuration, and SQLite companion files. Store this privately and retain the original machine association.')
h('Restore only through a controlled recovery')
step(7,'Keep the failed data tree for diagnosis. While the coordinator is stopped, have IT restore one complete matching backup, use the recorded software build, and keep the original hostname/port and keys. Do not overlay unrelated backups or combine a database with different secrets.')
step(8,'Have the maintainer verify database integrity and saved configuration. Start the coordinator, compare the public trust files/fingerprints, and repeat a two-client check before reopening the lab. Test restoration on an isolated spare machine; never run two reachable servers with the same identity at once.')
note('Keep data off the live sync path','Run KSAT from its installed local locations. Do not move the live database into Google Drive, OneDrive, or a network-synchronized folder. Copy completed, protected backups separately under institutional policy; do not add them to the distribution link.')

page('13 Update an existing lab and retain records')
h('Choose the scope of the update')
p('For this faculty-timer fix, update only the coordinator if clients already have the 10 September integrity update. Skip step 4 and run pilot checks with existing clients. New labs and older installations need both products.')
step(1,'Schedule downtime outside exams. Finish active attempts, resolve pending uploads, close finished assessments as appropriate, and record device counts, student/test counts, and representative results. Keep the current installers and checksum record.')
step(2,'Stop the coordinator and make the complete backup from page 13. Preserve client data before updating, especially PCs with retained attempts. Obtain the new verified coordinator and client Setup files from the same dated release.')
step(3,'Run the coordinator Setup as administrator on the server. Existing valid hostname/port settings are reused. The update replaces program files and retains ProgramData. Start the server and check the faculty page, bank/test records, and public trust.')
step(4,'Update two pilot clients first using client Setup as administrator. Existing valid client configuration is reused, and trust, device identity, attempts, and cached packs are retained. The installer controls the authority service during installation.')
step(5,'Repeat sign-in, a disposable assessment, integrity checks, and upload/review checks before updating the remaining PCs. Record hashes per PC: a displayed version of 2.0.0 cannot distinguish these test builds. Confirm every PC runs an approved compatible build.')
h('Earlier installations and migration stops')
p('Some early 2.0 clients have an older unauthenticated state format. Setup can deliberately stop and require a reviewed legacy-state migration. Keep the service stopped, preserve the full client data tree, and contact IT or the maintainer. After reviewing state counts and the backup, the documented installer option is /CONFIRMLEGACYSTATEMIGRATION=1. It is not a repair shortcut or a validation bypass.')
p('A legacy Aptitude Lab 1.x installation needs a separate controlled migration using a disposable copy first. Do not copy its database into this installation, overwrite an existing coordinator, or assume these client/server installers automatically migrate old data.')
h('Change the address or renew certificates')
p('Keep the same stable hostname and port when possible. If the server IP or network adapters change, fix DNS and have IT check the certificate: startup validates its recorded LAN addresses and can require controlled certificate renewal. Test before resuming. Do not edit saved JSON files or delete secrets to regenerate trust.')
p('For a deliberate address change, IT must use the documented service-stopped client --update-config --base-url workflow. It verifies the existing trust and refuses pending attempts. A new hostname must already be covered by the server certificate. A new coordinator CA requires a planned trust migration for every client.')
p('To renew an expired server certificate while preserving its CA, back up first, stop the coordinator, and run the following from Administrator PowerShell. Restart and retest afterwards:')
code('& "C:\Program Files\KSAT Coordinator\KSATCoordinator.exe" `\n  --renew-certificate --confirm-renewal')
note('Uninstall is not a reset','Uninstall removes program files and owned service/firewall/CA items as applicable, but deliberately retains ProgramData. Reinstalling does not erase the lab or automatically point it at another server. Keep retained data until an approved retention or recovery decision.')

page('14 Troubleshoot the network in order')
p('Work from the local service outward. Record the exact error and time. Repeat the same check on a known-good pilot client to distinguish a single-PC problem from a server or network outage.')
table(['Symptom','Checks and next action'],[
('Client local page will not open','Confirm http://127.0.0.1:8010 on that PC. Check KSATLabClientAuthority in services.msc. Verify the client installer completed; inspect its error/log before repairing. Do not open port 8010 to the LAN.'),
('Local page opens but server is unavailable','Check server power and that KSATCoordinator is actually running. Repeat ping by IP and hostname, then Test-NetConnection on the recorded port. Check the server did not reboot or sleep.'),
('IP ping passes and hostname ping fails','Correct DNS or the managed hosts mapping. Use the exact pretested hostname, not a guessed .local suffix. Verify that it resolves to the intended machine.'),
('TCP connection fails','Check coordinator process, chosen port, firewall profile/rule, VLAN routing, VPN route and network isolation. A working Internet connection or ping does not prove TCP 8443 is open.'),
('TCP passes but HTTPS validation fails','Check clocks, certificate expiry, exact URL/hostname, server CA trust, and matching metadata. A TCP pass is not a TLS pass. Do not switch to HTTP, use insecure flags, or click through certificate errors.'),
('Only another room or VLAN fails','Ask IT to verify inter-VLAN routing, ACLs and allowed source ranges. Do not broaden access to unmanaged networks. Repeat tests on a representative PC from each segment.'),
('Server works until restart or next day','Verify the DHCP reservation and DNS record. Start the coordinator shortcut again. Check sleep/update behavior and whether the active network profile changed.')],[2.08,4.79])
h('Useful read only diagnostics')
code('hostname\nipconfig /all\nGet-NetConnectionProfile\nping -4 -n 4 ksat-server.example.edu\nTest-NetConnection ksat-server.example.edu -Port 8443\nGet-Service KSATLabClientAuthority\nGet-Date')
p('Use the service command on a client. On the server, inspect the coordinator process and listening port using Task Manager or IT’s standard tools. Record any startup error dialog and relevant Windows Event Viewer entries; the packaged coordinator has no console window.')
note('Firewall ownership errors','If Setup reports a missing, duplicated, or changed owned firewall rule, preserve the installer log and ask IT/maintainer to reconcile the recorded configuration. Do not remove arbitrary rules or disable the firewall to force installation through.')

page('15 Troubleshoot setup and assessments')
table(['Symptom','Checks and next action'],[
('Hash mismatch or blocked download','Stop. Confirm the dated release with the sender and download again. Check institutional download policy. Never substitute an older 2.0.0 binary just because its filename matches.'),
('Invalid trust or configuration','Verify both public files came from this exact server, hashes match, JSON was not edited, and the HTTPS URL is correct. Existing installs retain their prior configuration; reinstall is not a server-change procedure.'),
('Coordinator already running or locked','Locate the existing coordinator process and use it. A lifecycle lock protects the data folder. Do not delete lock files or start two copies to bypass the protection.'),
('Device is inactive or unregistered','Check connection to the correct server. Inspect Managed lab computers on Faculty → Tests. Reactivate only after IT verifies the device identity. A reimaged machine normally needs a new identity.'),
('Wrong ID or password','Check Student versus Faculty role, Student ID/USN, case-sensitive password and whether the account exists on this department’s server. Do not delete an account with exam history to reset access.'),
('No assessment or Start not available','Verify the student is signed in, the correct server/test is used, content download is ready, Faculty launched the test, and the start window is still open. Select Refresh. Check prior completion or attempt state.'),
('Full screen or resume fails','Use the app’s entry/resume button in an ordinary Edge/Chrome tab. Check browser permissions, keyboard/mouse focus and institutional browser policy. F11 alone does not satisfy the app gate.'),
('Bank import fails or count differs','Check v2 package versus legacy pair, original ZIP structure, matching pair revision, missing assets and the exact validation error. Compare the sender’s expected counts; request a corrected package.'),
('Saved submission remains pending','Keep the same client and coordinator running. Restore network access and allow retries. A local score after Close can precede upload confirmation; verify the faculty accepted result.'),
('Faculty assistance or integrity error','Preserve the PC, attempt and error reference. Do not clear browser storage, edit SQLite, copy the attempt to another PC, delete state/anchors, or repeatedly reinstall. Escalate with diagnostic details.'),
('Review answers is unavailable','Faculty must explicitly Close the assessment. Start-window expiry alone is insufficient. Refresh and check the client connection and attempt state.'),
('A desktop switch is not flagged','Record Windows/browser versions, action and timing; repeat in a disposable exam. Browser-only detection has limits. Retain the limitation in acceptance records and seek kiosk/native monitoring if required.')],[2.07,4.8])
h('Information to send to support')
p('Provide department, PC names, Windows/browser versions, release checksums, exact time/time zone, steps, visible error/reference, assessment and attempt IDs, service status, ping/TCP results, and whether other PCs work. Setup logs are normally under the installing user’s %TEMP% as Setup Log files. Record their path when an installer fails.')
p('Redact passwords, session cookies/tokens, private keys, student personal data and bank answers from ordinary screenshots or email. Preserve full diagnostic copies only through the institution’s protected support channel.')

page('16 Build identity and technical references')
p('These hashes identify the integrity and faculty timer test binaries dated 10 September 2026. Source revision: d6509f72c93dfa024c09e6d7419e648e8d5a8f1f. All filenames retain version 2.0.0. A later release needs its own checksum record and guide review.')
h('Files used for installation')
p('KSATCoordinatorSetup-2.0.0.exe  •  28,140,152 bytes',True)
code('23d8d4a39a9b766feb5fa5eb388c0955a82a364a4ec6393e0b7e25bd36cf5eba')
p('KSATClientSetup-2.0.0.exe  •  28,045,800 bytes',True)
code('51d05b7e8279e363f7bf1ae015b2662f339f9cdcb42edfb8ea77be39f7eb31c6')
h('Standalone executables for reference')
p('KSATCoordinator-2.0.0.exe  •  26,468,288 bytes',True)
code('74a99936608a858d443f87c173fab26d95c245be60b83b79a48de653802255b9')
p('KSATClient-2.0.0.exe  •  26,373,752 bytes',True)
code('1b54acacf00433c407b90710fc5bdcb341e899f9cf106f78f58e855f11619566')
p('Executable signing identity: CN=KSAT TEST SIGNING IDENTITY - NOT FOR PRODUCTION. A checksum verifies the downloaded bytes against this record; it does not make a test signing identity production trusted.')
h('Installed locations at default settings')
table(['Purpose','Location or name'],[
('Server application','C:\Program Files\KSAT Coordinator\KSATCoordinator.exe'),
('Server complete data and backup source','C:\ProgramData\KSAT Coordinator'),
('Client application','C:\Program Files\KSAT Client\KSATClient.exe'),
('Client protected data','C:\ProgramData\KSAT Client'),
('Client Windows service','KSATLabClientAuthority'),
('Server public trust export','Server data folder\public'),
('Server local faculty page','https://127.0.0.1:8443'),
('Student local client page','http://127.0.0.1:8010')],[2.1,4.77])
h('Sources for maintainers')
p('Application instructions were checked against the current installer scripts, entry points, faculty/client UI, and project runbooks. Use the pinned revision below when comparing this guide with software behavior.')
link('KSAT source at this guide revision','https://github.com/raviasha/Aptitude-Test/tree/d6509f72c93dfa024c09e6d7419e648e8d5a8f1f')
link('Project operations runbook','https://github.com/raviasha/Aptitude-Test/blob/d6509f72c93dfa024c09e6d7419e648e8d5a8f1f/docs/distributed-assessment-operations.md')
link('Microsoft Get FileHash reference','https://learn.microsoft.com/en-us/powershell/module/microsoft.powershell.utility/get-filehash')
link('Microsoft certificate store administration reference','https://learn.microsoft.com/en-us/windows-server/administration/windows-commands/certutil')

page('17 Department handover and acceptance record')
p('Complete this sheet and retain it with the dated installers, bank inventory, network worksheet, and protected backup record. Approval here is for the agreed test deployment only; the build remains NOT FOR PRODUCTION.')
table(['Item','Recorded value'],[
('Department and lab','__________________________________________________'),
('IT owner and faculty operator','__________________________________________________'),
('Install date and guide revision','__________________________________________________'),
('Server hostname and reserved IP','__________________________________________________'),
('Coordinator URL and port','__________________________________________________'),
('Pilot PCs and Windows/browser versions','__________________________________________________'),
('Number of installed client seats','__________________________________________________'),
('Distribution and checksum record location','__________________________________________________'),
('CA file SHA256 record location','__________________________________________________'),
('Public metadata SHA256 record location','__________________________________________________'),
('Bank filenames revisions and counts record','__________________________________________________'),
('Complete backup location and restore owner','__________________________________________________')],[2.55,4.32])
h('Acceptance checklist')
for t in [
'[  ] Two clients passed preinstallation IP and hostname checks, or documented IT ICMP exceptions and routing verification.',
'[  ] Both clients passed post-install TCP, trust, enrollment, and ordinary-user sign-in checks.',
'[  ] Server startup after reboot and automatic client service startup were tested.',
'[  ] Firewall/VLAN access is restricted to managed lab PCs; an unmanaged network cannot reach the coordinator.',
'[  ] Bank counts, sample questions, diagrams, correct answers, and solutions were verified.',
'[  ] Start, full screen, autosave, submission, faculty CSV, Close, review, and sign-out passed.',
'[  ] Tab/window/desktop/lock tests and network outage recovery were recorded, including any missed events.',
'[  ] All planned seats use the same approved client/server build; class-size rehearsal completed.',
'[  ] Complete backup exists; restore owner, startup operator, and support contact are assigned.'
]:p(t)
p('Outstanding issues and restrictions: __________________________________________')
p('________________________________________________________________________')
p('IT acceptance and date: __________________   Faculty acceptance and date: __________________')

OUT.mkdir(parents=True,exist_ok=True)
for parent in [doc.styles.element, doc.element]:
    for border in list(parent.xpath('.//w:pBdr')):
        border.getparent().remove(border)
path=OUT/(NAME+'.docx')
doc.save(path)
print(path)
