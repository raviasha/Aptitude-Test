from pathlib import Path
import hashlib
import json
import zipfile
from docx import Document
from docx.shared import Inches

out=Path(__file__).resolve().parents[1]
reference=out/'KSAT_Quick_Installation_Guide.docx'
reference_digest=hashlib.sha256(reference.read_bytes()).hexdigest()
doc=Document(reference)
for element in list(doc.element.body):
    if element.tag.endswith('}sectPr'):
        continue
    doc.element.body.remove(element)
doc.core_properties.title='KSAT Quick User Guide'
doc.core_properties.subject='Essential faculty and student steps for the installed client and server edition'

def p(text='',bold=False):
    para=doc.add_paragraph()
    para.add_run(text).bold=bold
    return para

def h(text): doc.add_heading(text,1)

def step(number,text):
    para=p()
    para.paragraph_format.left_indent=Inches(.18)
    para.paragraph_format.first_line_indent=Inches(-.18)
    para.add_run(f'{number}. ').bold=True
    para.add_run(text)

doc.add_paragraph('KSAT Quick User Guide','Title')
p('Faculty and students  |  TEST BUILD - NOT FOR PRODUCTION',True)
p('Use this guide after the lab has been installed. Faculty instructions are on this page; student instructions are on page 2.')

h('Faculty steps to prepare and launch an assessment')
step(1,'Open KSAT Faculty Coordinator on the server. Select Faculty and your Department, then sign in using the faculty credentials. Keep the Coordinator running and the server awake throughout the session.')
step(2,'If the bank is not already available, open Question banks, select its original ZIP under Upload a package, and choose Upload v2 package. For a supplied HTML/JSON pair, use Upload a pair. Confirm that the bank appears.')
step(3,'Open Tests. Enter a Test name, select the Question bank and Difficulty, and enter the required number of questions for each chapter. Check the total and select Create test. The initial duration is one minute per question.')
step(4,'Wait for Content to show Ready. Once students are signed in and ready, select Launch. Students have ten minutes to start; each student’s exam timer begins when their attempt starts.')

h('Monitor timing and add extra time')
p('In Test library, Exam time left shows the remaining time for active students. A range means their deadlines differ. Start window shows how long students have to begin; it is separate from their exam time.')
step(1,'To add time for the whole assessment, select Extend beside the test and enter a reason. Each successful use adds five minutes to active timers and future starts. Check Exam time left and min added on Faculty. This does not reopen or extend the start window.')
step(2,'For one student, select Inspect / manage attempts. Enter the assessment ID shown, then that student’s attempt ID. Enter extend as the action, give a reason, and enter the additional minutes. Confirm the updated timing on Faculty and the student PC.')

h('Finish the assessment and keep results')
step(1,'Open Overview to check Submitted results, scores, and Exam integrity records. Select Download results CSV to save the results. Investigate any reported violation with the invigilator.')
step(2,'After students have finished, select Close beside the test when you are ready to release correct answers and solutions. Students can then use Review answers. Expiry of the start window alone does not release answers.')
step(3,'For another class or sitting, use Duplicate to create a fresh copy of the test. Existing attempts and results stay with the original test. Keep the server running until pending submissions are received.')
p('Sign-in problem: open Students and use Release login for the affected student if an old login is blocking them. An ongoing attempt must continue on the same PC. Do not delete a student or test to fix a login problem; deletion removes their related records.')

doc.add_page_break()
h('Student steps to sign in and start')
step(1,'Open the KSAT Lab Client shortcut on your assigned PC. If you do not have an account, select New student? Create account. Enter your correct name, Student ID / USN, Class and Section, and a password of at least six characters. Then select your Department and sign in.')
step(2,'Wait for faculty to ask you to begin. Find the correct assessment and select Start, or Download if shown. Download also starts the attempt after the files arrive. If the assessment is missing, select Refresh and ask faculty.')
step(3,'Enter full screen when prompted. Use Enter full screen if the screen asks you to return. Check that the assessment and timer are visible. Stay on this PC for the whole attempt.')

h('Answer questions')
step(1,'Select an answer and check the save message. Answers save automatically on your PC. Use Previous and Next to move between questions. You can change answers while the attempt is still open.')
step(2,'Keep the exam visible, focused, and in full screen. Switching tabs, windows or desktops, minimizing, leaving full screen, or copying and pasting can record violations. If a return-to-full-screen message appears, return to the exam and select Enter full screen.')
p('The timer continues during interruptions and while you are away from the exam. Tell the invigilator promptly if anything stops you working.')

h('Submit and view your result')
step(1,'Check your answers before selecting Submit assessment. Submission locks the answers; they cannot be changed afterwards. When time runs out, the attempt also locks and is submitted automatically.')
step(2,'Wait for the result to be confirmed by the server. If Upload pending appears, your saved answers will upload automatically. Keep the PC running and tell faculty. A message asking for Faculty assistance also needs their attention.')
step(3,'Correct answers and solutions become available after faculty closes the assessment. Select Review answers when available. You can also find finished tests under Completed assessments. Waiting for Faculty to close means the review has not yet been released.')
step(4,'Use Back to assessments to return to the list. When submission is confirmed and Sign out is available, sign out before leaving the PC or letting another student use it.')

h('If something goes wrong')
p('Tell the invigilator the message on screen and your Student ID / USN. If the network drops but the assessment still works, continue answering on the same PC. Leave the client open while the connection recovers. Do not switch computers, create another account, or clear saved browser data to restart an attempt.')

path=out/'KSAT_Quick_User_Guide.docx'
doc.save(path)
assert hashlib.sha256(reference.read_bytes()).hexdigest()==reference_digest
with zipfile.ZipFile(reference) as before, zipfile.ZipFile(path) as after:
    assert set(before.namelist())==set(after.namelist())
    changed=[name for name in before.namelist() if before.read(name)!=after.read(name)]
    assert set(changed)<= {'word/document.xml','docProps/core.xml'},changed
(out/'guide-work'/'user-guide-fidelity.json').write_text(json.dumps({'reference_sha256':reference_digest,'changed_parts':changed},indent=2),encoding='utf-8')
print(path)
